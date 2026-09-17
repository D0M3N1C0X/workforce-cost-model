"""The Excel model reproduces the pandas model, formula by formula.

Evaluating the full workbook takes minutes in pure Python, so this test builds a sample
workbook and calculates it with the `formulas` package. CI also recalculates the full
workbook with LibreOffice (see .github/workflows/ci.yml).
"""
import pytest

formulas = pytest.importorskip("formulas")

import build_workbook  # noqa: E402
import check_workbook  # noqa: E402
import model  # noqa: E402
from deterministic import same_content  # noqa: E402


@pytest.fixture(scope="module")
def sample_model(tmp_path_factory):
    path = tmp_path_factory.mktemp("wb") / "sample.xlsx"
    m = build_workbook.build(model.run(check_workbook.sample(2, seed=11)), path)
    return m, path


def test_every_reconciliation_row_matches_and_nothing_errors(sample_model, capsys):
    m, path = sample_model
    values = check_workbook.values_from_formulas(path)
    assert check_workbook.check(values, m.wb.sheetnames, len(m.checks)) == 0, capsys.readouterr().out


def test_a_wrong_python_value_is_caught(sample_model, tmp_path):
    from openpyxl import load_workbook
    m, path = sample_model
    wb = load_workbook(path)
    ws = wb["Reconciliation"]
    row = next(r for r in range(7, 7 + len(m.checks)) if isinstance(ws[f"D{r}"].value, float))
    ws[f"D{row}"].value *= 1 + 1e-6
    broken = tmp_path / "broken.xlsx"
    wb.save(broken)
    values = check_workbook.values_from_formulas(broken)
    assert check_workbook.check(values, wb.sheetnames, len(m.checks)) == 1


def test_every_calculating_sheet_is_reconciled(sample_model):
    m, _ = sample_model
    areas = {area for area, *_ in m.checks}
    assert set(m.wb.sheetnames) - areas == {"Cover", "Reconciliation"}


def test_the_workbook_is_reproducible(out, tmp_path):
    a, b = tmp_path / "a.xlsx", tmp_path / "b.xlsx"
    build_workbook.build(out, a)
    build_workbook.build(out, b)
    assert same_content(a.read_bytes(), b.read_bytes())
    assert a.read_bytes() == b.read_bytes()
