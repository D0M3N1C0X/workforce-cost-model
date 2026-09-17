"""
Proves the workbook's formulas reproduce the pandas pipeline.

openpyxl writes formulas without results, so a spreadsheet engine has to calculate them
first. Two engines are supported:

    # LibreOffice (what CI uses): recalculate a copy, then check the saved values
    python src/recalc_libreoffice.py deliverables/workforce_cost_model.xlsx build/
    python src/check_workbook.py build/workforce_cost_model.xlsx

    # the pure-Python `formulas` package, on a small sample (quick to run locally)
    python src/check_workbook.py --sample 2    # 2 stayers, leavers and joiners per department

The check fails on any formula error anywhere in the workbook, and on any row of the
Reconciliation sheet that does not match.
"""
import argparse
import sys
from pathlib import Path

ERRORS = ("#DIV/0!", "#N/A", "#NAME?", "#NULL!", "#NUM!", "#REF!", "#VALUE!", "Err:")


def values_from_saved(path: Path) -> dict[tuple[str, str], object]:
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True, read_only=True)
    out = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    out[(ws.title, cell.coordinate)] = cell.value
    return out


def values_from_formulas(path: Path) -> dict[tuple[str, str], object]:
    import formulas
    solution = formulas.ExcelModel().loads(str(path)).finish().calculate()
    out = {}
    for key, ranges in solution.items():
        if "!" not in key:
            continue
        sheet, ref = key.rsplit("!", 1)
        sheet = sheet.strip("'").split("]", 1)[-1]
        if ":" in ref:
            continue
        value = ranges.value[0][0] if hasattr(ranges, "value") else ranges
        if hasattr(value, "item"):
            value = value.item()
        out[(sheet.upper(), ref.replace("$", ""))] = value
    return out


def check(values: dict, sheet_names: list[str], n_checks: int) -> int:
    upper = {s.upper(): s for s in sheet_names}
    get = lambda sheet, ref: values.get((sheet, ref), values.get((sheet.upper(), ref)))

    errors = [(s, r, v) for (s, r), v in values.items()
              if not isinstance(v, (int, float)) and str(v).startswith(ERRORS)]
    for s, r, v in errors[:20]:
        print(f"::error::formula error {v} at {upper.get(s, s)}!{r}")

    mismatches = []
    for row in range(7, 7 + n_checks):
        if get("Reconciliation", f"G{row}") != "Yes":
            mismatches.append((get("Reconciliation", f"A{row}"), get("Reconciliation", f"B{row}"),
                               get("Reconciliation", f"D{row}"), get("Reconciliation", f"E{row}")))
    for m in mismatches[:30]:
        print("::error::mismatch", m)

    print(f"formula errors: {len(errors)}")
    print(f"reconciliation: {n_checks - len(mismatches)} of {n_checks} match")
    print(f"summary cell: {get('Reconciliation', 'B3')}")
    return 1 if errors or mismatches or n_checks == 0 else 0


def sample(per_cell: int, seed: int):
    """A few stayers, leavers and joiners in every entity and department, so every formula has data."""
    import model
    r = model.roster()
    parts = []
    for (c, d), g in r.groupby(["country", "department"]):
        stay = g[(g["hire_date"] <= model.OPENING) & (g["exit_date"].isna())]
        left = g[g["exit_date"].notna()]
        join = g[g["hire_date"] > model.OPENING]
        for part in (stay, left, join):
            parts.append(part.sample(n=min(len(part), per_cell), random_state=seed))
    import pandas as pd
    return pd.concat(parts).sort_values("employee_id")


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("workbook", nargs="?", type=Path, help="a workbook already recalculated by LibreOffice or Excel")
    ap.add_argument("--sample", type=int, help="build a sample workbook (N stayers, leavers and joiners per department) and evaluate it with `formulas`")
    ap.add_argument("--seed", type=int, default=7)
    ap.add_argument("--formulas", action="store_true", help="calculate the given workbook with `formulas` (slow at full size)")
    args = ap.parse_args()

    if args.sample:
        import build_workbook
        import model
        from config import ROOT
        path = ROOT / "build" / "sample_model.xlsx"
        m = build_workbook.build(model.run(sample(args.sample, args.seed)), path)
        print(f"sample workbook: {len(m.r)} employees, {len(m.checks)} checks")
        return check(values_from_formulas(path), m.wb.sheetnames, len(m.checks))

    if not args.workbook:
        ap.error("give a recalculated workbook, or --sample N")
    from openpyxl import load_workbook
    names = load_workbook(args.workbook, read_only=True).sheetnames
    values = values_from_formulas(args.workbook) if args.formulas else values_from_saved(args.workbook)
    # one row per check from row 7, each with an item in column B
    n = sum(1 for (s, r) in values if s.upper() == "RECONCILIATION" and r.startswith("B")
            and r[1:].isdigit() and int(r[1:]) >= 7)
    return check(values, names, n)


if __name__ == "__main__":
    sys.exit(main())
