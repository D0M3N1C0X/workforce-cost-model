"""
Builds deliverables/workforce_cost_model.xlsx: the people-cost model a finance team would use.

Every number is a live formula over the Roster and Settings sheets: change a contribution rate,
an attrition driver or a scenario and the budget, the variance and the forecast recalculate.
The Reconciliation sheet holds the pandas results and checks each one against its formula.

Conventions as in pay-transparency-readiness-kit: blue text = input, black = formula, green =
link to another sheet, yellow fill = key assumption, Arial throughout, no dynamic arrays.
"""
import math
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as col
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.hyperlink import Hyperlink

import config as C
import model
from deterministic import normalise

OUTPUT = C.DELIVERABLES / "workforce_cost_model.xlsx"
REPO = "https://github.com/D0M3N1C0X/workforce-cost-model"

FONT = "Arial"
BLUE, GREEN, INK, MUTED, WHITE = "0000FF", "008000", "1B2430", "5F6B7A", "FFFFFF"
HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
SUB_FILL = PatternFill("solid", fgColor="3E5C76")
BAND_FILL = PatternFill("solid", fgColor="EEF2F7")
YELLOW = PatternFill("solid", fgColor="FFFF00")
PCT, PCT1, COUNT, FTE = "0.00%", "0.0%", "#,##0", "#,##0.0"
EUR = '"€"#,##0;-"€"#,##0;"-"'
EURK = '"€"#,##0,"k";-"€"#,##0,"k";"-"'
DATE = "d mmm yyyy"
MONTH = "mmm yy"


def font(color=INK, bold=False, italic=False, size=10):
    return Font(name=FONT, color=color, bold=bold, italic=italic, size=size)


def put(ws, ref, value, *, color=INK, bold=False, italic=False, size=10, fmt=None, fill=None, wrap=False, align=None):
    cell = ws[ref]
    cell.value = value
    cell.font = font(color, bold, italic, size)
    if fmt:
        cell.number_format = fmt
    if fill:
        cell.fill = fill
    if wrap or align:
        cell.alignment = Alignment(wrap_text=wrap, horizontal=align, vertical="top")
    return cell


def header(ws, row, labels, start=1, height=30, fill=HEADER_FILL):
    for i, label in enumerate(labels):
        c = ws.cell(row=row, column=start + i, value=label)
        c.font = font(WHITE, bold=True)
        c.fill = fill
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[row].height = height


def title(ws, text, subtitle=None):
    put(ws, "A1", text, bold=True, size=14)
    if subtitle:
        put(ws, "A2", subtitle, color=MUTED, italic=True)


def widths(ws, spec):
    for k, v in spec.items():
        ws.column_dimensions[k].width = v


def plain(x):
    if x is None:
        return None
    if hasattr(x, "item"):
        x = x.item()
    if isinstance(x, float) and math.isnan(x):
        return None
    return x


class Model:
    def __init__(self, out: dict):
        self.o = out
        self.r = out["roster"]
        self.wb = Workbook()
        self.refs = {}
        self.checks = []
        self.cells = model.cells()
        self.sheet_of = {s[0]: f"FY27 {s[0]}" for s in C.SCENARIOS}

    def check(self, area, item, value, ref):
        self.checks.append((area, item, plain(value), ref))

    def rng(self, name):
        letter = self.rcol[name]
        return f"Roster!${letter}$2:${letter}${self.rlast}"

    # ------------------------------------------------------------------------------------------
    def build(self, path: Path):
        wb = self.wb
        names = ["Cover", "Dashboard", "Scenarios", "Variance FY26", "Actuals FY26", "Budget FY26",
                 "Drivers FY27", *self.sheet_of.values(), "Roster", "Settings", "Reconciliation"]
        wb.active.title = names[0]
        for n in names[1:]:
            wb.create_sheet(n)
        self.settings(wb["Settings"])
        self.roster(wb["Roster"])
        self.actuals(wb["Actuals FY26"])
        self.budget(wb["Budget FY26"])
        self.variance(wb["Variance FY26"])
        self.drivers(wb["Drivers FY27"])
        for i, s in enumerate(C.SCENARIOS):
            self.forecast(wb[self.sheet_of[s[0]]], i, s)
        self.scenarios(wb["Scenarios"])
        self.dashboard(wb["Dashboard"])
        self.reconciliation(wb["Reconciliation"])
        self.cover(wb["Cover"], names)
        for ws in wb.worksheets:
            ws.sheet_view.showGridLines = False
            ws.page_setup.orientation = "landscape"
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0
            ws.sheet_properties.pageSetUpPr.fitToPage = True
        wb.calculation.fullCalcOnLoad = True
        stamp = datetime(2026, 1, 1)
        wb.properties.creator = "Domenico Perroni"
        wb.properties.title = "Workforce Cost Model"
        wb.properties.created = wb.properties.modified = stamp
        path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(path)
        normalise(path)

    # ------------------------------------------------------------------------------------------
    def settings(self, ws):
        title(ws, "Settings", "Every assumption. Blue cells are inputs; yellow ones move the results most.")
        widths(ws, {"A": 3, "B": 52, "C": 14, "D": 14, "E": 14, "F": 14, "G": 14, "H": 70})
        r = 4

        def section(text):
            nonlocal r
            r += 1
            put(ws, f"B{r}", text, bold=True, size=11)
            r += 1

        def line(key, label, value, fmt, note="", key_assumption=False, formula=False):
            nonlocal r
            put(ws, f"B{r}", label)
            put(ws, f"C{r}", value, color=INK if formula else BLUE, fmt=fmt, fill=YELLOW if key_assumption else None)
            put(ws, f"H{r}", note, color=MUTED, italic=True)
            self.refs[key] = f"Settings!$C${r}"
            r += 1

        section("Calendar")
        line("fy26_start", "First day of FY2026", C.FY2026_START, DATE, "Fiscal years run July to June.")
        line("opening", "Opening date (FY2026 budget population)", f"=C{r - 1}-1", DATE, formula=True)
        line("closing", "Closing date (FY2027 starting population)", f"=EOMONTH(C{r - 2},{C.MONTHS - 1})", DATE,
             formula=True)

        section("Italy: employer costs")
        line("it_ivs", "INPS pension contribution (IVS), employer share", C.IT_IVS, PCT,
             "33% in total, 23.81% paid by the employer.")
        line("it_other", "Other INPS contributions and INAIL premium", C.IT_OTHER, PCT,
             "Illustrative: depends on the INPS classification and INAIL risk rate. Set per company.", True)
        line("tfr_div", "TFR divisor (annual base pay / divisor)", C.TFR_DIVISOR, "0.0",
             "Art. 2120 Civil Code. Applied to base pay only.")

        section("Poland: employer costs (2026)")
        line("pl_pen", "Pension contribution (emerytalne), capped", C.PL_PENSION, PCT)
        line("pl_dis", "Disability contribution (rentowe), capped", C.PL_DISABILITY, PCT)
        line("pl_acc", "Accident contribution (wypadkowe)", C.PL_ACCIDENT, PCT,
             "Reference rate. A payer with ten or more insured people has its own risk-based rate: use it.")
        line("pl_fp", "Labour Fund (Fundusz Pracy), not capped", C.PL_LABOUR_FUND, PCT)
        line("pl_fgsp", "Guaranteed Employee Benefits Fund (FGŚP), not capped", C.PL_FGSP, PCT)
        line("pl_ppk", "PPK employer basic contribution, not capped", C.PL_PPK, PCT)
        line("ppk_part", "Share of pay covered by PPK participants", C.PPK_PARTICIPATION, PCT1,
             "Illustrative.", True)
        line("cap_pln", "Annual cap on the pension and disability base, PLN", C.PL_CAP_PLN, COUNT,
             "2026: 30 x 9,420 PLN projected average wage. Update when the 2027 limit is published.")
        line("fx", "EUR/PLN", C.EUR_PLN, "0.0000", "ECB reference rate, monthly average June 2026.", True)
        line("cap_eur", "Annual cap in euro", f"=C{r - 2}/C{r - 1}", EUR, formula=True)
        line("min_pln", "Monthly minimum wage 2026, PLN", C.PL_MIN_WAGE_PLN, COUNT,
             "No Polish salary may fall below it: checked on the Reconciliation sheet.")
        line("min_eur", "Annual minimum wage in euro (full time)", f"=C{r - 1}*12/{self.refs['fx'].split('!')[1]}",
             EUR, formula=True)

        section("Target bonus by level (paid at target)")
        header(ws, r, ["Level", "Target"], start=2, height=18)
        r += 1
        first = r
        for level, t in C.BONUS_TARGET.items():
            put(ws, f"B{r}", level, color=BLUE)
            put(ws, f"C{r}", t, color=BLUE, fmt=PCT1)
            r += 1
        self.refs["levels"] = f"Settings!$B${first}:$B${r - 1}"
        self.refs["targets"] = f"Settings!$C${first}:$C${r - 1}"

        section("Growth by department (net, over the year)")
        header(ws, r, ["Department", "FY2026 budget", "FY2027 forecast"], start=2, height=18)
        r += 1
        first = r
        for d in C.DEPARTMENTS:
            put(ws, f"B{r}", d, color=BLUE)
            put(ws, f"C{r}", C.BUDGET_GROWTH[d], color=BLUE, fmt=PCT1)
            put(ws, f"D{r}", C.FORECAST_GROWTH[d], color=BLUE, fmt=PCT1, fill=YELLOW)
            r += 1
        self.refs["depts"] = f"Settings!$B${first}:$B${r - 1}"
        self.refs["growth26"] = f"Settings!$C${first}:$C${r - 1}"
        self.refs["growth27"] = f"Settings!$D${first}:$D${r - 1}"

        section("FY2027 by entity")
        header(ws, r, ["Entity", "Pay review", "Cost per hire", "Remediation (€ a year)", "On-cost factor"],
               start=2, height=30)
        r += 1
        first = r
        for c in C.ENTITIES:
            put(ws, f"B{r}", c, color=BLUE)
            put(ws, f"C{r}", C.MERIT[c], color=BLUE, fmt=PCT1, fill=YELLOW)
            put(ws, f"D{r}", C.COST_PER_HIRE[c], color=BLUE, fmt=EUR)
            put(ws, f"E{r}", C.REMEDIATION[c], color=BLUE, fmt=EUR)
            if c == "IT":
                f = f"=1+{self.refs['it_ivs']}+{self.refs['it_other']}+1/{self.refs['tfr_div']}"
            else:
                f = (f"=1+{self.refs['pl_pen']}+{self.refs['pl_dis']}+{self.refs['pl_acc']}+{self.refs['pl_fp']}"
                     f"+{self.refs['pl_fgsp']}+{self.refs['pl_ppk']}*{self.refs['ppk_part']}")
            put(ws, f"F{r}", f, fmt="0.0000")
            self.check("Settings", f"On-cost factor {c}", model.loaded_factor(c), f"Settings!F{r}")
            r += 1
        for key, letter in (("ent", "B"), ("merit", "C"), ("cph", "D"), ("rem", "E"), ("factor", "F")):
            self.refs[key] = f"Settings!${letter}${first}:${letter}${r - 1}"
        put(ws, f"H{first}", f"Remediation: upper-bound cost of closing flagged pay gaps, from "
                             f"pay-transparency-readiness-kit at commit {C.KIT_COMMIT}. Cost per hire is illustrative.",
            color=MUTED, italic=True)
        line("merit_month", "Pay review effective from month (1 = July)", C.MERIT_MONTH, "0")
        line("tth", "Time to hire, months", C.TIME_TO_HIRE, "0")

        section("Scenarios")
        header(ws, r, ["Scenario", "Attrition x", "Replacement", "Growth x", "Freeze from month",
                       "Remediation from month"], start=2, height=30)
        r += 1
        first = r
        for s in C.SCENARIOS:
            put(ws, f"B{r}", s[0], color=BLUE)
            for j, (v, fmt) in enumerate(zip(s[1:], ["0.00", PCT1, "0.00", "0", "0"])):
                put(ws, f"{col(3 + j)}{r}", v, color=BLUE, fmt=fmt, fill=YELLOW)
            r += 1
        self.refs["scen_first"] = first
        self.refs["scen_names"] = f"Settings!$B${first}:$B${r - 1}"
        put(ws, f"H{first}", "Freeze and remediation months: 0 means never. Month 4 = October 2026, month 7 = January 2027.",
            color=MUTED, italic=True)
        r += 1
        put(ws, f"B{r}", "Scenario shown on the Dashboard", bold=True)
        put(ws, f"C{r}", C.SCENARIOS[0][0], color=BLUE, fill=YELLOW)
        dv = DataValidation(type="list", formula1=f"={self.refs['scen_names']}")
        ws.add_data_validation(dv)
        dv.add(f"C{r}")
        self.refs["selected"] = f"Settings!$C${r}"

    # ------------------------------------------------------------------------------------------
    def roster(self, ws):
        r = self.r
        inputs = ["employee_id", "country", "department", "job_level", "fte", "salary", "hire_date", "exit_date"]
        derived = ["exit_key", "base", "bonus", "gross", "capped", "employer_social", "tfr", "loaded_cost"]
        names = inputs + derived
        self.rcol = {n: col(i) for i, n in enumerate(names, start=1)}
        self.rlast = len(r) + 1
        header(ws, 1, names, height=32)
        for n in derived:
            ws[f"{self.rcol[n]}1"].fill = SUB_FILL
        c = self.rcol
        R = self.refs
        blue, black = font(BLUE), font()
        for i, row in enumerate(r[inputs].itertuples(index=False), start=2):
            for n, v in zip(inputs, row):
                cell = ws[f"{c[n]}{i}"]
                if n in ("hire_date", "exit_date"):
                    cell.value = None if pd.isna(v) else v.to_pydatetime().date()
                    cell.number_format = DATE
                else:
                    cell.value = plain(v)
                cell.font = blue
            L = {k: f"{v}{i}" for k, v in c.items()}
            f = {
                "exit_key": f'=IF({L["exit_date"]}="",DATE(2099,12,31),{L["exit_date"]})',
                "base": f"={L['salary']}*{L['fte']}",
                "bonus": f"={L['base']}*INDEX({R['targets']},MATCH({L['job_level']},{R['levels']},0))",
                "gross": f"={L['base']}+{L['bonus']}",
                "capped": f"=MIN({L['gross']},{R['cap_eur']})",
                "employer_social": (f'=IF({L["country"]}="IT",{L["gross"]}*{R["it_ivs"]}+{L["gross"]}*{R["it_other"]},'
                                    f'{L["capped"]}*{R["pl_pen"]}+{L["capped"]}*{R["pl_dis"]}+{L["gross"]}*{R["pl_acc"]}'
                                    f'+{L["gross"]}*{R["pl_fp"]}+{L["gross"]}*{R["pl_fgsp"]}'
                                    f'+{L["gross"]}*{R["pl_ppk"]}*{R["ppk_part"]})'),
                "tfr": f'=IF({L["country"]}="IT",{L["base"]}/{R["tfr_div"]},0)',
                "loaded_cost": f"={L['gross']}+{L['employer_social']}+{L['tfr']}",
            }
            fmts = {"exit_key": DATE, "base": EUR, "bonus": EUR, "gross": EUR, "capped": EUR,
                    "employer_social": EUR, "tfr": EUR, "loaded_cost": EUR}
            for n, formula in f.items():
                cell = ws[f"{c[n]}{i}"]
                cell.value = formula
                cell.font = black
                cell.number_format = fmts[n]
        for n, letter in c.items():
            ws.column_dimensions[letter].width = 13
        ws.column_dimensions[c["department"]].width = 18
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = f"A1:{col(len(names))}{self.rlast}"
        # a sample of employees for the reconciliation
        for i in list(range(0, len(r), max(1, len(r) // 40)))[:40]:
            self.check("Roster", f"{r.loc[i, 'employee_id']} loaded cost", r.loc[i, "loaded_cost"],
                       f"Roster!{c['loaded_cost']}{i + 2}")
        self.check("Roster", "Total loaded cost, all employees", r["loaded_cost"].sum(),
                   f"SUM(Roster!{c['loaded_cost']}2:{c['loaded_cost']}{self.rlast})")
        below = int(((r["country"] == "PL") & (r["salary"] < model.pl_min_wage_eur())).sum())
        self.check("Roster", "Polish full-time salaries below the 2026 minimum wage", below,
                   f'SUMPRODUCT(({self.rng("country")}="PL")*({self.rng("salary")}<{R["min_eur"]}))')

    def when(self, ref):
        """SUMIFS criteria for 'on the payroll at date ref'."""
        return f'{self.rng("hire_date")},"<="&{ref},{self.rng("exit_key")},">"&{ref}'

    def cell_criteria(self, row):
        return f'{self.rng("country")},$A{row},{self.rng("department")},$B{row}'

    # ------------------------------------------------------------------------------------------
    def actuals(self, ws):
        title(ws, "FY2026 actuals: month-end FTE and fully loaded cost",
              "Counted from the Roster: on the payroll at each month end. Cost is the annual loaded cost / 12.")
        widths(ws, {"A": 8, "B": 18, **{col(k): 10 for k in range(3, 30)}})
        months = [col(k) for k in range(3, 15)]           # C..N FTE
        cmonths = [col(k) for k in range(17, 29)]         # Q..AB cost
        put(ws, "A4", "Month end", bold=True)
        for j, (a, b) in enumerate(zip(months, cmonths)):
            put(ws, f"{a}4", f"=EOMONTH({self.refs['fy26_start']},{j})", fmt=MONTH, bold=True, align="right")
            put(ws, f"{b}4", f"={a}4", fmt=MONTH, bold=True, align="right")
        header(ws, 5, ["Entity", "Department"] + ["FTE"] * 12 + ["Average FTE", ""] + ["Cost"] * 12 + ["Total cost"])
        act = self.o["actuals"].set_index(["country", "department", "month"])
        self.act_rows = {}
        row = 6
        for cty, d in self.cells:
            self.act_rows[(cty, d)] = row
            put(ws, f"A{row}", cty, color=BLUE)
            put(ws, f"B{row}", d, color=BLUE)
            for j, (a, b) in enumerate(zip(months, cmonths), start=1):
                put(ws, f"{a}{row}", f"=SUMIFS({self.rng('fte')},{self.cell_criteria(row)},{self.when(f'{a}$4')})", fmt=FTE)
                put(ws, f"{b}{row}", f"=SUMIFS({self.rng('loaded_cost')},{self.cell_criteria(row)},{self.when(f'{a}$4')})/12",
                    fmt=EURK)
                if j in (1, 6, 12):
                    self.check("Actuals FY26", f"{cty} {d} month {j} FTE", act.loc[(cty, d, j), "fte"], f"'Actuals FY26'!{a}{row}")
                    self.check("Actuals FY26", f"{cty} {d} month {j} cost", act.loc[(cty, d, j), "cost"], f"'Actuals FY26'!{b}{row}")
            put(ws, f"O{row}", f"=AVERAGE(C{row}:N{row})", fmt=FTE, bold=True)
            put(ws, f"AC{row}", f"=SUM(Q{row}:AB{row})", fmt=EUR, bold=True)
            row += 1
        self.act_total = {}
        for cty in C.ENTITIES:
            rows = [self.act_rows[(cty, d)] for d in C.DEPARTMENTS]
            put(ws, f"A{row}", cty, bold=True)
            put(ws, f"B{row}", "Total", bold=True)
            for a in months + ["O"] + cmonths + ["AC"]:
                put(ws, f"{a}{row}", f"=SUM({a}{rows[0]}:{a}{rows[-1]})", bold=True,
                    fmt=FTE if a in months + ["O"] else (EUR if a == "AC" else EURK), fill=BAND_FILL)
            self.act_total[cty] = row
            row += 1
        ws.freeze_panes = "C6"

    def budget(self, ws):
        title(ws, "FY2026 budget: the plan set at 30 June 2025",
              "Opening FTE and cost per FTE from the Roster at the opening date; FTE grows in a straight line to the planned level.")
        widths(ws, {"A": 8, "B": 18, "C": 11, "D": 13, "E": 10, **{col(k): 10 for k in range(6, 34)}})
        mcols = [col(k) for k in range(6, 18)]        # F..Q FTE
        ccols = [col(k) for k in range(20, 32)]       # T..AE cost
        put(ws, "E4", "Month", bold=True)
        for j, (a, b) in enumerate(zip(mcols, ccols), start=1):
            put(ws, f"{a}4", j, bold=True, align="right")
            put(ws, f"{b}4", j, bold=True, align="right")
        header(ws, 5, ["Entity", "Department", "Opening FTE", "Cost per FTE", "Growth"] + ["FTE"] * 12
               + ["Average FTE", ""] + ["Cost"] * 12 + ["Total cost"])
        bud = self.o["budget"].set_index(["country", "department", "month"]).sort_index()
        self.bud_rows = {}
        row = 6
        opening = self.refs["opening"]
        for cty, d in self.cells:
            self.bud_rows[(cty, d)] = row
            put(ws, f"A{row}", cty, color=BLUE)
            put(ws, f"B{row}", d, color=BLUE)
            put(ws, f"C{row}", f"=SUMIFS({self.rng('fte')},{self.cell_criteria(row)},{self.when(opening)})", fmt=FTE)
            put(ws, f"D{row}", f"=SUMIFS({self.rng('loaded_cost')},{self.cell_criteria(row)},{self.when(opening)})/C{row}",
                fmt=EUR)
            put(ws, f"E{row}", f"=INDEX({self.refs['growth26']},MATCH(B{row},{self.refs['depts']},0))", fmt=PCT1)
            for a, b in zip(mcols, ccols):
                put(ws, f"{a}{row}", f"=$C{row}*(1+$E{row}*{a}$4/12)", fmt=FTE)
                put(ws, f"{b}{row}", f"={a}{row}*$D{row}/12", fmt=EURK)
            put(ws, f"R{row}", f"=AVERAGE(F{row}:Q{row})", fmt=FTE, bold=True)
            put(ws, f"AF{row}", f"=SUM(T{row}:AE{row})", fmt=EUR, bold=True)
            self.check("Budget FY26", f"{cty} {d} month 12 FTE", bud.loc[(cty, d, 12), "fte"], f"'Budget FY26'!Q{row}")
            self.check("Budget FY26", f"{cty} {d} total cost", bud.loc[(cty, d), "cost"].sum(), f"'Budget FY26'!AF{row}")
            row += 1
        ws.freeze_panes = "C6"

    def variance(self, ws):
        title(ws, "FY2026: actual against budget",
              "Volume = FTE difference at budget cost per FTE. Rate = cost per FTE difference at actual FTE. "
              "Mix = department volumes minus the entity's pure volume effect.")
        widths(ws, {"A": 8, "B": 18, **{col(k): 13 for k in range(3, 14)}})
        heads = ["Entity", "Department", "Budget FTE", "Actual FTE", "Budget cost", "Actual cost", "Budget cost per FTE",
                 "Actual cost per FTE", "Volume", "Mix", "Rate", "Variance"]
        header(ws, 4, heads, height=32)
        vd = self.o["variance_dept"].set_index(["country", "department"])
        ve = self.o["variance_entity"].set_index("country")
        row = 5
        self.var_rows = {}
        for cty in C.ENTITIES:
            first = row
            for d in C.DEPARTMENTS:
                a, b = self.act_rows[(cty, d)], self.bud_rows[(cty, d)]
                put(ws, f"A{row}", cty, color=BLUE)
                put(ws, f"B{row}", d, color=BLUE)
                put(ws, f"C{row}", f"='Budget FY26'!R{b}", color=GREEN, fmt=FTE)
                put(ws, f"D{row}", f"='Actuals FY26'!O{a}", color=GREEN, fmt=FTE)
                put(ws, f"E{row}", f"='Budget FY26'!AF{b}", color=GREEN, fmt=EUR)
                put(ws, f"F{row}", f"='Actuals FY26'!AC{a}", color=GREEN, fmt=EUR)
                put(ws, f"G{row}", f"=E{row}/C{row}", fmt=EUR)
                put(ws, f"H{row}", f"=F{row}/D{row}", fmt=EUR)
                put(ws, f"I{row}", f"=(D{row}-C{row})*G{row}", fmt=EUR)
                put(ws, f"K{row}", f"=(H{row}-G{row})*D{row}", fmt=EUR)
                put(ws, f"L{row}", f"=F{row}-E{row}", fmt=EUR, bold=True)
                v = vd.loc[(cty, d)]
                for letter, key in (("I", "volume"), ("K", "rate"), ("L", "variance")):
                    self.check("Variance FY26", f"{cty} {d} {key}", v[key], f"'Variance FY26'!{letter}{row}")
                row += 1
            last = row - 1
            put(ws, f"A{row}", cty, bold=True)
            put(ws, f"B{row}", "Total", bold=True)
            for letter in "CDEF":
                put(ws, f"{letter}{row}", f"=SUM({letter}{first}:{letter}{last})", bold=True,
                    fmt=FTE if letter in "CD" else EUR, fill=BAND_FILL)
            put(ws, f"G{row}", f"=E{row}/C{row}", fmt=EUR, bold=True, fill=BAND_FILL)
            put(ws, f"H{row}", f"=F{row}/D{row}", fmt=EUR, bold=True, fill=BAND_FILL)
            put(ws, f"I{row}", f"=(D{row}-C{row})*G{row}", fmt=EUR, bold=True, fill=BAND_FILL)
            put(ws, f"J{row}", f"=SUM(I{first}:I{last})-I{row}", fmt=EUR, bold=True, fill=BAND_FILL)
            put(ws, f"K{row}", f"=SUM(K{first}:K{last})", fmt=EUR, bold=True, fill=BAND_FILL)
            put(ws, f"L{row}", f"=F{row}-E{row}", fmt=EUR, bold=True, fill=BAND_FILL)
            e = ve.loc[cty]
            for letter, key in (("I", "volume"), ("J", "mix"), ("K", "rate"), ("L", "variance"),
                                ("E", "budget_cost"), ("F", "actual_cost")):
                self.check("Variance FY26", f"{cty} total {key}", e[key], f"'Variance FY26'!{letter}{row}")
            self.var_rows[cty] = row
            row += 2
        put(ws, f"A{row}", "Check: volume + mix + rate = variance for each entity.", color=MUTED, italic=True)
        for cty, vr in self.var_rows.items():
            row += 1
            put(ws, f"B{row}", cty)
            put(ws, f"C{row}", f'=IF(ABS(I{vr}+J{vr}+K{vr}-L{vr})<0.01,"Adds up","Does not add up")')

    def drivers(self, ws):
        title(ws, "FY2027 drivers, from FY2026 actuals",
              "Starting FTE and cost per FTE at 30 June 2026; attrition is FY2026 leavers' FTE over average FTE.")
        widths(ws, {"A": 8, "B": 18, **{col(k): 13 for k in range(3, 11)}})
        header(ws, 4, ["Entity", "Department", "Starting FTE", "Cost per FTE", "FY2026 leavers (FTE)",
                       "FY2026 average FTE", "Attrition, annual", "Growth FY2027", "Pay review", "Cost per hire"],
               height=32)
        attr = self.o["attrition"].set_index(["country", "department"])
        start = self.o["start"].set_index(["country", "department"])
        closing, opening = self.refs["closing"], self.refs["opening"]
        self.drv_rows = {}
        for row, (cty, d) in enumerate(self.cells, start=5):
            self.drv_rows[(cty, d)] = row
            put(ws, f"A{row}", cty, color=BLUE)
            put(ws, f"B{row}", d, color=BLUE)
            put(ws, f"C{row}", f"=SUMIFS({self.rng('fte')},{self.cell_criteria(row)},{self.when(closing)})", fmt=FTE)
            put(ws, f"D{row}", f"=SUMIFS({self.rng('loaded_cost')},{self.cell_criteria(row)},{self.when(closing)})/C{row}",
                fmt=EUR)
            put(ws, f"E{row}", (f'=SUMIFS({self.rng("fte")},{self.cell_criteria(row)},{self.rng("exit_key")},">"&{opening},'
                                f'{self.rng("exit_key")},"<="&{closing})'), fmt=FTE)
            put(ws, f"F{row}", f"='Actuals FY26'!O{self.act_rows[(cty, d)]}", color=GREEN, fmt=FTE)
            put(ws, f"G{row}", f"=E{row}/F{row}", fmt=PCT1)
            put(ws, f"H{row}", f"=INDEX({self.refs['growth27']},MATCH(B{row},{self.refs['depts']},0))", fmt=PCT1)
            put(ws, f"I{row}", f"=INDEX({self.refs['merit']},MATCH(A{row},{self.refs['ent']},0))", fmt=PCT1)
            put(ws, f"J{row}", f"=INDEX({self.refs['cph']},MATCH(A{row},{self.refs['ent']},0))", fmt=EUR)
            self.check("Drivers FY27", f"{cty} {d} starting FTE", start.loc[(cty, d), "fte"], f"'Drivers FY27'!C{row}")
            self.check("Drivers FY27", f"{cty} {d} cost per FTE", start.loc[(cty, d), "rate"], f"'Drivers FY27'!D{row}")
            self.check("Drivers FY27", f"{cty} {d} attrition", attr.loc[(cty, d), "rate"], f"'Drivers FY27'!G{row}")

    def forecast(self, ws, k, scenario):
        name = scenario[0]
        s_row = self.refs["scen_first"] + k
        title(ws, f"FY2027 forecast: {name}",
              "FTE(m) = FTE(m-1) - leavers + hires. Leavers = FTE(m-1) x monthly attrition. Hires = backfills of leavers "
              "from 'time to hire' months earlier x replacement, plus planned growth; none after a hiring freeze.")
        widths(ws, {"A": 8, "B": 18, "C": 12, "D": 11, **{col(k2): 10 for k2 in range(5, 22)}})
        params = ["Attrition x", "Replacement", "Growth x", "Freeze from month", "Remediation from month"]
        for j, p in enumerate(params):
            put(ws, f"{col(5 + 2 * j)}3", p, bold=True)
            put(ws, f"{col(6 + 2 * j)}3", f"=Settings!{col(3 + j)}{s_row}", color=GREEN)
        att_x, repl, growth_x, freeze, rem_from = (f"${col(6 + 2 * j)}$3" for j in range(5))
        months = [col(k2) for k2 in range(5, 17)]          # E..P
        put(ws, "D5", 0, bold=True, align="right")
        for j, a in enumerate(months, start=1):
            put(ws, f"{a}5", j, bold=True, align="right")
        header(ws, 6, ["Entity", "Department", "Line", "Opening"] + [f"M{j}" for j in range(1, 13)]
               + ["FY total / avg", "Monthly attrition", "Monthly growth", "Cost per FTE", "Pay review"])
        f, rem = self.o["forecasts"][name]
        fi = f.set_index(["country", "department", "month"]).sort_index()
        tth, mm = self.refs["tth"], self.refs["merit_month"]
        row = 7
        blocks = {c: [] for c in C.ENTITIES}
        for cty, d in self.cells:
            drv = self.drv_rows[(cty, d)]
            fr, lr, hr, cr, kr = row, row + 1, row + 2, row + 3, row + 4
            blocks[cty].append((fr, cr, kr, hr))
            for rr, label in zip((fr, lr, hr, cr, kr), ("FTE", "Leavers", "Hires", "Payroll cost", "Hiring cost")):
                put(ws, f"A{rr}", cty, color=BLUE)
                put(ws, f"B{rr}", d, color=BLUE)
                put(ws, f"C{rr}", label, bold=label == "FTE")
            put(ws, f"D{fr}", f"='Drivers FY27'!C{drv}", color=GREEN, fmt=FTE)
            put(ws, f"R{fr}", f"='Drivers FY27'!G{drv}*{att_x}/12", fmt="0.000%")
            put(ws, f"S{fr}", f"='Drivers FY27'!H{drv}*{growth_x}", fmt=PCT1)
            put(ws, f"T{fr}", f"='Drivers FY27'!D{drv}", color=GREEN, fmt=EUR)
            put(ws, f"U{fr}", f"='Drivers FY27'!I{drv}", color=GREEN, fmt=PCT1)
            prev = "D"
            for a in months:
                # Backfills look only at earlier months' leavers: a range reaching the current or later
                # months would make Excel see a circular reference through INDEX.
                earlier = "E" if prev == "D" else prev
                put(ws, f"{a}{lr}", f"={prev}{fr}*$R{fr}", fmt=FTE)
                put(ws, f"{a}{hr}", (f"=IF(AND({freeze}>0,{a}$5>={freeze}),0,"
                                     f"IF({a}$5>{tth},INDEX($E{lr}:{earlier}{lr},{a}$5-{tth}),$D{fr}*$R{fr})*{repl}"
                                     f"+$D{fr}*$S{fr}/12)"), fmt=FTE)
                put(ws, f"{a}{fr}", f"={prev}{fr}-{a}{lr}+{a}{hr}", fmt=FTE, bold=True)
                put(ws, f"{a}{cr}", f"={a}{fr}*($T{fr}*(1+IF({a}$5>={mm},$U{fr},0)))/12", fmt=EURK)
                put(ws, f"{a}{kr}", f"={a}{hr}*'Drivers FY27'!$J${drv}", fmt=EURK)
                prev = a
            put(ws, f"Q{fr}", f"=AVERAGE(E{fr}:P{fr})", fmt=FTE, bold=True)
            for rr, fmt in ((lr, FTE), (hr, FTE), (cr, EUR), (kr, EUR)):
                put(ws, f"Q{rr}", f"=SUM(E{rr}:P{rr})", fmt=fmt, bold=True)
            for j in (1, 7, 12):
                self.check(f"FY27 {name}", f"{cty} {d} month {j} FTE", fi.loc[(cty, d, j), "fte"], f"'{ws.title}'!{months[j - 1]}{fr}")
            self.check(f"FY27 {name}", f"{cty} {d} payroll cost", fi.loc[(cty, d), "cost"].sum(), f"'{ws.title}'!Q{cr}")
            self.check(f"FY27 {name}", f"{cty} {d} hires", fi.loc[(cty, d), "hires"].sum(), f"'{ws.title}'!Q{hr}")
            for rr in (fr, lr, hr, cr, kr):
                if (rr - 7) // 5 % 2 == 0:
                    for a in ["A", "B", "C", "D"] + months + ["Q"]:
                        ws[f"{a}{rr}"].fill = BAND_FILL
            row += 5
        # entity totals
        row += 1
        self.fc_tot = getattr(self, "fc_tot", {})
        for cty in C.ENTITIES:
            fr_list = [b[0] for b in blocks[cty]]
            cr_list = [b[1] for b in blocks[cty]]
            kr_list = [b[2] for b in blocks[cty]]
            hr_list = [b[3] for b in blocks[cty]]
            ent_row = f"MATCH(\"{cty}\",{self.refs['ent']},0)"
            tf, tc, tr, tk, th = row, row + 1, row + 2, row + 3, row + 4
            for rr, label in zip((tf, tc, tr, tk, th), ("Total FTE", "Payroll cost", "Pay-equity remediation",
                                                         "Hiring cost", "Hires")):
                put(ws, f"A{rr}", cty, bold=True)
                put(ws, f"C{rr}", label, bold=True)
            for a in ["D"] + months:
                put(ws, f"{a}{tf}", "=" + "+".join(f"{a}{x}" for x in fr_list), fmt=FTE, bold=True, fill=BAND_FILL)
            for a in months:
                put(ws, f"{a}{tc}", "=" + "+".join(f"{a}{x}" for x in cr_list), fmt=EURK, fill=BAND_FILL)
                put(ws, f"{a}{tr}", (f"=IF(AND({rem_from}>0,{a}$5>={rem_from}),"
                                     f"INDEX({self.refs['rem']},{ent_row})*INDEX({self.refs['factor']},{ent_row})/12,0)"),
                    fmt=EURK, fill=BAND_FILL)
                put(ws, f"{a}{tk}", "=" + "+".join(f"{a}{x}" for x in kr_list), fmt=EURK, fill=BAND_FILL)
                put(ws, f"{a}{th}", "=" + "+".join(f"{a}{x}" for x in hr_list), fmt=FTE, fill=BAND_FILL)
            put(ws, f"Q{tf}", f"=AVERAGE(E{tf}:P{tf})", fmt=FTE, bold=True, fill=BAND_FILL)
            for rr in (tc, tr, tk, th):
                put(ws, f"Q{rr}", f"=SUM(E{rr}:P{rr})", fmt=EUR if rr != th else FTE, bold=True, fill=BAND_FILL)
            self.fc_tot[(name, cty)] = {"fte_avg": f"'{ws.title}'!Q{tf}", "closing": f"'{ws.title}'!P{tf}",
                                        "payroll": f"'{ws.title}'!Q{tc}", "remediation": f"'{ws.title}'!Q{tr}",
                                        "hiring": f"'{ws.title}'!Q{tk}", "hires": f"'{ws.title}'!Q{th}",
                                        "sheet": ws.title, "fte_row": tf}
            row += 6
        ws.freeze_panes = "E7"

    def scenarios(self, ws):
        title(ws, "FY2027 scenarios against FY2026 actuals",
              "People cost = payroll + pay-equity remediation. Hiring cost is shown apart. "
              "Bridge: volume at FY2026 cost per FTE, then the change in cost per FTE at FY2027 FTE.")
        widths(ws, {"A": 20, "B": 8, **{col(k): 14 for k in range(3, 16)}})
        heads = ["Scenario", "Entity", "FY2026 average FTE", "FY2026 cost", "FY2027 average FTE", "FY2027 closing FTE",
                 "FY2027 payroll", "Remediation", "FY2027 people cost", "Change vs FY2026", "Bridge: volume",
                 "Bridge: rate", "Hires", "Hiring cost", "Key"]
        header(ws, 4, heads, height=32)
        summ = self.o["summary"].set_index(["scenario", "country"])
        row = 5
        self.scen_rows = {}
        for s in C.SCENARIOS:
            name = s[0]
            for cty in C.ENTITIES:
                t = self.fc_tot[(name, cty)]
                at = self.act_total[cty]
                put(ws, f"A{row}", name, color=BLUE)
                put(ws, f"B{row}", cty, color=BLUE)
                put(ws, f"C{row}", f"='Actuals FY26'!O{at}", color=GREEN, fmt=FTE)
                put(ws, f"D{row}", f"='Actuals FY26'!AC{at}", color=GREEN, fmt=EUR)
                put(ws, f"E{row}", f"={t['fte_avg']}", color=GREEN, fmt=FTE)
                put(ws, f"F{row}", f"={t['closing']}", color=GREEN, fmt=FTE)
                put(ws, f"G{row}", f"={t['payroll']}", color=GREEN, fmt=EUR)
                put(ws, f"H{row}", f"={t['remediation']}", color=GREEN, fmt=EUR)
                put(ws, f"I{row}", f"=G{row}+H{row}", fmt=EUR, bold=True)
                put(ws, f"J{row}", f"=I{row}/D{row}-1", fmt=PCT1)
                put(ws, f"K{row}", f"=(E{row}-C{row})*(D{row}/C{row})", fmt=EUR)
                put(ws, f"L{row}", f"=(G{row}/E{row}-D{row}/C{row})*E{row}", fmt=EUR)
                put(ws, f"M{row}", f"={t['hires']}", color=GREEN, fmt=FTE)
                put(ws, f"N{row}", f"={t['hiring']}", color=GREEN, fmt=EUR)
                put(ws, f"O{row}", f'=A{row}&"|"&B{row}')
                sm = summ.loc[(name, cty)]
                for letter, key in (("C", "fy26_fte"), ("D", "fy26_cost"), ("E", "fy27_fte"), ("F", "closing_fte"),
                                    ("G", "fy27_payroll"), ("H", "fy27_remediation"), ("I", "fy27_cost"),
                                    ("K", "bridge_volume"), ("L", "bridge_rate"), ("M", "fy27_hires"),
                                    ("N", "fy27_hiring_cost")):
                    self.check("Scenarios", f"{name} {cty} {key}", sm[key], f"Scenarios!{letter}{row}")
                self.scen_rows[(name, cty)] = row
                row += 1
        self.scen_range = (5, row - 1)
        row += 1
        put(ws, f"A{row}", "Check: FY2026 cost + volume + rate + remediation = FY2027 people cost.", color=MUTED, italic=True)
        first, last = self.scen_range
        put(ws, f"A{row + 1}", (f'=IF(SUMPRODUCT(ABS(D{first}:D{last}+K{first}:K{last}+L{first}:L{last}'
                                f'+H{first}:H{last}-I{first}:I{last}))<0.1,"Every bridge adds up","A bridge does not add up")'))

    def dashboard(self, ws):
        title(ws, "Workforce cost: dashboard",
              "Pick a scenario in Settings (the yellow cell under Scenarios); every figure below follows it.")
        widths(ws, {"A": 3, "B": 34, "C": 16, "D": 16, "E": 16, "F": 3, **{col(k): 12 for k in range(7, 16)}})
        put(ws, "B4", "Scenario", bold=True)
        put(ws, "C4", f"={self.refs['selected']}", color=GREEN, bold=True, size=12)
        header(ws, 6, ["", "Measure", *C.ENTITIES, "Both"], start=1, height=20)
        first, last = self.scen_range
        key = lambda c: f'"|"&"{c}"'
        lookup = lambda colname, c: (f"=INDEX(Scenarios!${colname}${first}:${colname}${last},"
                                     f"MATCH($C$4&{key(c)},Scenarios!$O${first}:$O${last},0))")
        rows = [("FY2026 actual cost", "D", EUR), ("FY2027 people cost", "I", EUR), ("Change vs FY2026", "J", PCT1),
                ("FY2027 average FTE", "E", FTE), ("FY2027 closing FTE", "F", FTE), ("Pay-equity remediation", "H", EUR),
                ("Hires", "M", FTE), ("Hiring cost", "N", EUR)]
        r = 7
        shown = self.o["summary"].query("scenario == @C.SCENARIOS[0][0]").sum(numeric_only=True)
        both = {"D": shown["fy26_cost"], "I": shown["fy27_cost"], "E": shown["fy27_fte"],
                "F": shown["closing_fte"], "H": shown["fy27_remediation"], "M": shown["fy27_hires"],
                "N": shown["fy27_hiring_cost"]}
        for label, colname, fmt in rows:
            put(ws, f"B{r}", label, bold=label.startswith("FY2027 people"))
            for j, c in enumerate(C.ENTITIES):
                put(ws, f"{col(3 + j)}{r}", lookup(colname, c), color=GREEN, fmt=fmt)
            if colname == "J":
                put(ws, f"E{r}", f"=E{r - 1}/E{r - 2}-1", fmt=fmt)
                value = shown["fy27_cost"] / shown["fy26_cost"] - 1
            else:
                put(ws, f"E{r}", f"=C{r}+D{r}", fmt=fmt, bold=True)
                value = both[colname]
            self.check("Dashboard", f"{label}, both entities ({C.SCENARIOS[0][0]} selected)", value, f"Dashboard!E{r}")
            r += 1
        r += 1
        put(ws, f"B{r}", "FY2026 against budget", bold=True, size=11)
        r += 1
        header(ws, r, ["", "", *C.ENTITIES, "Both"], start=1, height=18)
        r += 1
        for label, letter in (("Budget", "E"), ("Actual", "F"), ("Volume", "I"), ("Mix", "J"), ("Rate", "K"),
                              ("Variance", "L")):
            put(ws, f"B{r}", label, bold=label == "Variance")
            for j, c in enumerate(C.ENTITIES):
                put(ws, f"{col(3 + j)}{r}", f"='Variance FY26'!{letter}{self.var_rows[c]}", color=GREEN, fmt=EUR)
            put(ws, f"E{r}", f"=C{r}+D{r}", fmt=EUR, bold=label == "Variance")
            key = {"Budget": "budget_cost", "Actual": "actual_cost"}.get(label, label.lower())
            self.check("Dashboard", f"FY2026 {label.lower()}, both entities", self.o["variance_entity"][key].sum(),
                       f"Dashboard!E{r}")
            r += 1

        # chart data: total FTE, FY2026 actual then FY2027 per scenario
        r += 2
        data_top = r
        put(ws, f"B{r}", "Chart data: total FTE, both entities", bold=True)
        r += 1
        heads = ["Month", "FY2026 actual"] + [s[0] for s in C.SCENARIOS]
        for j, h in enumerate(heads):
            put(ws, f"{col(2 + j)}{r}", h, bold=True)
        hdr = r
        r += 1
        for m in range(1, 25):
            put(ws, f"B{r}", f"=EOMONTH({self.refs['fy26_start']},{m - 1})", fmt=MONTH)
            if m <= 12:
                a = col(2 + m)
                put(ws, f"C{r}", f"='Actuals FY26'!{a}{self.act_total['IT']}+'Actuals FY26'!{a}{self.act_total['PL']}", fmt=FTE)
            for j, s in enumerate(C.SCENARIOS):
                t_it, t_pl = self.fc_tot[(s[0], "IT")], self.fc_tot[(s[0], "PL")]
                a = col(m - 8) if m > 12 else None     # FY2027 month m-12 sits in column E + (m-13)
                if m == 12:
                    put(ws, f"{col(4 + j)}{r}", f"=C{r}", fmt=FTE)
                elif m > 12:
                    put(ws, f"{col(4 + j)}{r}",
                        f"='{t_it['sheet']}'!{a}{t_it['fte_row']}+'{t_pl['sheet']}'!{a}{t_pl['fte_row']}", fmt=FTE)
                    if m == 24:
                        f, _ = self.o["forecasts"][s[0]]
                        self.check("Dashboard", f"Chart: {s[0]} FTE in June 2027",
                                   f.loc[f["month"] == C.MONTHS, "fte"].sum(), f"Dashboard!{col(4 + j)}{r}")
            r += 1
        line = LineChart()
        line.title = "Total FTE: FY2026 actual and FY2027 scenarios"
        line.y_axis.title = "FTE"
        line.y_axis.majorGridlines = None
        line.height, line.width = 8, 17
        line.add_data(Reference(ws, min_col=3, max_col=3 + len(C.SCENARIOS), min_row=hdr, max_row=r - 1),
                      titles_from_data=True)
        line.set_categories(Reference(ws, min_col=2, min_row=hdr + 1, max_row=r - 1))
        colours = ["1B2430", "2A78D6", "EB6834", "1BAF7A", "EDA100"]
        for s_, c_ in zip(line.series, colours):
            s_.graphicalProperties.line.solidFill = c_
            s_.graphicalProperties.line.width = 22000
            s_.smooth = False
        ws.add_chart(line, "G4")

        bar = BarChart()
        bar.type = "col"
        bar.title = "FY2027 people cost by scenario"
        bar.y_axis.majorGridlines = None
        bar.height, bar.width = 8, 17
        # compact table for the bar chart
        r += 1
        put(ws, f"B{r}", "Chart data: FY2027 people cost", bold=True)
        r += 1
        put(ws, f"B{r}", "Scenario", bold=True)
        for j, c in enumerate(C.ENTITIES):
            put(ws, f"{col(3 + j)}{r}", c, bold=True)
        bh = r
        for s in C.SCENARIOS:
            r += 1
            put(ws, f"B{r}", s[0])
            for j, c in enumerate(C.ENTITIES):
                put(ws, f"{col(3 + j)}{r}", f"=Scenarios!I{self.scen_rows[(s[0], c)]}", color=GREEN, fmt=EUR)
        bar.add_data(Reference(ws, min_col=3, max_col=4, min_row=bh, max_row=r), titles_from_data=True)
        bar.set_categories(Reference(ws, min_col=2, min_row=bh + 1, max_row=r))
        for s_, c_ in zip(bar.series, ["2A78D6", "EB6834"]):
            s_.graphicalProperties.solidFill = c_
        ws.add_chart(bar, "G22")

    def reconciliation(self, ws):
        title(ws, "Reconciliation: workbook formulas against the Python pipeline",
              "Column D was written by src/build_workbook.py from the pandas results; column E is the live formula.")
        widths(ws, {"A": 16, "B": 58, "C": 3, "D": 18, "E": 18, "F": 14, "G": 8})
        header(ws, 6, ["Area", "Item", "", "Python value", "Workbook value", "Difference", "Match"])
        first, last = 7, 6 + len(self.checks)
        put(ws, "B3", (f'=IF(COUNTIF(G{first}:G{last},"No")=0,"All "&COUNTA(B{first}:B{last})&" checks match",'
                       f'COUNTIF(G{first}:G{last},"No")&" of "&COUNTA(B{first}:B{last})&" checks do not match")'),
            bold=True, size=12)
        for r, (area, item, value, ref) in enumerate(self.checks, start=first):
            put(ws, f"A{r}", area)
            put(ws, f"B{r}", item)
            fmt = "0.000000" if isinstance(value, float) else None
            put(ws, f"D{r}", value, color=BLUE, fmt=fmt)
            put(ws, f"E{r}", f"={ref}", color=GREEN, fmt=fmt)
            put(ws, f"F{r}", f'=IF(AND(ISNUMBER(D{r}),ISNUMBER(E{r})),E{r}-D{r},"")', fmt="0.0E+00")
            put(ws, f"G{r}", (f'=IF(AND(D{r}="",E{r}=""),"Yes",IF(AND(ISNUMBER(D{r}),ISNUMBER(E{r})),'
                              f'IF(ABS(E{r}-D{r})<=1E-9*MAX(1,ABS(D{r})),"Yes","No"),IF(D{r}=E{r},"Yes","No")))'))
        for text, colour in (("Yes", "D5ECDC"), ("No", "F5C6C2")):
            ws.conditional_formatting.add(f"G{first}:G{last}", CellIsRule(operator="equal", formula=[f'"{text}"'],
                                                                         fill=PatternFill("solid", fgColor=colour)))
        ws.freeze_panes = "A7"
        self.recon_result = "Reconciliation!B3"

    def cover(self, ws, names):
        widths(ws, {"A": 3, "B": 24, "C": 100})
        put(ws, "B2", "Workforce Cost Model", bold=True, size=18)
        put(ws, "B3", "Italy and Poland: FY2026 budget against actual, and FY2027 people cost under four scenarios",
            color=MUTED, size=12)
        put(ws, "B5", "What it does", bold=True, size=11)
        lines = [
            "Prices every employee at fully loaded cost: base pay, target bonus, Italian INPS and TFR, Polish ZUS with the 2026 cap.",
            "Rebuilds FY2026 month-end FTE and cost from hire and exit dates, and splits the budget variance into volume, mix and rate.",
            "Forecasts FY2027 month by month from FY2026 attrition, backfills, planned growth and the January pay review.",
            "Compares four scenarios, including a hiring freeze and closing pay gaps before the first pay-transparency report.",
        ]
        for i, text in enumerate(lines, start=6):
            put(ws, f"C{i}", text)
        put(ws, "B11", "How to read it", bold=True, size=11)
        for i, (label, meaning, colour, fill) in enumerate([
            ("Blue text", "an input", BLUE, None), ("Black text", "a formula", INK, None),
            ("Green text", "a link to another sheet", GREEN, None),
            ("Yellow fill", "a key assumption, including the scenario shown", INK, YELLOW)], start=12):
            put(ws, f"B{i}", label, color=colour, fill=fill)
            put(ws, f"C{i}", meaning)
        put(ws, "B17", "Sheets", bold=True, size=11)
        purpose = {
            "Dashboard": "Headline figures and charts for the scenario selected in Settings.",
            "Scenarios": "Every scenario side by side, with the bridge from FY2026.",
            "Variance FY26": "Budget against actual by department: volume, mix and rate.",
            "Actuals FY26": "Month-end FTE and cost counted from the Roster.",
            "Budget FY26": "The plan: opening FTE, cost per FTE and planned growth.",
            "Drivers FY27": "Starting FTE, cost per FTE and attrition taken from FY2026.",
            "Roster": "One row per employee: inputs in blue, loaded cost in black.",
            "Settings": "Rates, caps, growth, pay review, cost per hire and the scenario table.",
            "Reconciliation": "Each figure checked against the Python pipeline.",
        }
        for s in C.SCENARIOS:
            purpose[self.sheet_of[s[0]]] = f"FY2027 month by month under the '{s[0]}' scenario."
        for i, n in enumerate(names[1:], start=18):
            cell = put(ws, f"B{i}", n, color="1F5FA8")
            cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{n}'!A1")
            put(ws, f"C{i}", purpose[n])
        r = 18 + len(names)
        put(ws, f"B{r}", "Check", bold=True, size=11)
        put(ws, f"C{r}", f"={self.recon_result}", color=GREEN, bold=True)
        put(ws, f"B{r + 2}", "Data", bold=True, size=11)
        put(ws, f"C{r + 2}", "Synthetic: the Italian and Polish entities of the organisation analysed in hr-people-analytics. "
                             "Budget, growth, pay review and cost per hire are invented for the demonstration.", wrap=True)
        ws.row_dimensions[r + 2].height = 28
        put(ws, f"B{r + 3}", "Limits", bold=True, size=11)
        put(ws, f"C{r + 3}", "Month-end population; salaries as at 30 June 2026 for all months; one cap and one exchange "
                             "rate for the whole year. Not payroll or tax advice: see docs/verification.md.", wrap=True)
        ws.row_dimensions[r + 3].height = 28
        put(ws, f"B{r + 5}", "Domenico Perroni", color=MUTED)
        cell = put(ws, f"C{r + 5}", REPO, color="1F5FA8")
        cell.hyperlink = REPO


def build(out: dict | None = None, path: Path = OUTPUT) -> Model:
    out = model.run() if out is None else out
    m = Model(out)
    m.build(path)
    return m


def main() -> None:
    m = build()
    print(f"workbook: {len(m.checks)} reconciliation checks -> {OUTPUT.relative_to(OUTPUT.parents[1])}")


if __name__ == "__main__":
    main()
